# -*- coding: utf-8 -*-
"""md2docx: 把本仓库的中文 Markdown 教程转成排版规整的 Word(.docx)。
支持: 标题层级、GFM 表格、代码围栏、有序/无序列表(含缩进)、引用块、
      行内 **加粗** / `代码` / [文字](链接)；中文字体使用宋体 + 标题微软雅黑。
用法: python tools/md2docx.py 文件1.md 文件2.md ...   (输出同目录同名 .docx)
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

EA_BODY, EA_HEAD, EN_BODY, EN_CODE = "宋体", "微软雅黑", "Times New Roman", "Consolas"
HEAD_COLOR = RGBColor(0x1F, 0x39, 0x64)
QUOTE_COLOR = RGBColor(0x55, 0x55, 0x55)
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")

def set_font(run, ea=EA_BODY, en=EN_BODY, size=11, bold=False, color=None, italic=False):
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    if color is not None: run.font.color.rgb = color
    run.font.name = en
    rPr = run._element.get_or_add_rPr(); rf = rPr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), en); rf.set(qn("w:hAnsi"), en); rf.set(qn("w:eastAsia"), ea)

def _shd(fill):
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); return shd

def shade_par(par, fill):
    par._p.get_or_add_pPr().append(_shd(fill))

def shade_cell(cell, fill):
    cell._tc.get_or_add_tcPr().append(_shd(fill))

def add_inline(par, text, size=11, base_ea=EA_BODY, color=None, bold=False):
    """递归处理行内 **粗体** / `代码` / [文字](链接)，支持粗体内再嵌套链接/代码。"""
    for tok in INLINE.split(text):
        if not tok: continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) >= 4:
            add_inline(par, tok[2:-2], size, base_ea, color, bold=True)   # 递归, 内部继承加粗
        elif tok.startswith("`") and tok.endswith("`") and len(tok) >= 2:
            set_font(par.add_run(tok[1:-1]), en=EN_CODE, ea=EN_CODE, size=size-0.5,
                     bold=bold, color=RGBColor(0xB0,0x30,0x50))
        elif tok.startswith("[") and "](" in tok:
            label = tok[1:tok.index("]")].replace(chr(96),"").replace("*","")
            set_font(par.add_run(label), ea=base_ea, size=size, bold=bold, color=RGBColor(0x11,0x55,0xCC))
        else:
            set_font(par.add_run(tok), ea=base_ea, size=size, bold=bold, color=color)

def flush_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    if len(cells) < 2: return
    head, body = cells[0], cells[2:]  # 第 1 行表头, 第 2 行分隔
    n = len(head)
    t = doc.add_table(rows=1, cols=n); t.style = "Table Grid"; t.autofit = True
    for j, h in enumerate(head):
        cp = t.rows[0].cells[j].paragraphs[0]; cp.paragraph_format.line_spacing = 1.1
        add_inline(cp, h, size=10, base_ea=EA_HEAD)
        for r in cp.runs: r.bold = True
        shade_cell(t.rows[0].cells[j], "E8EDF7")
    for row in body:
        cs = t.add_row().cells
        for j in range(n):
            txt = row[j] if j < len(row) else ""
            cp = cs[j].paragraphs[0]; cp.paragraph_format.line_spacing = 1.1
            add_inline(cp, txt, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def convert(md_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.size = Pt(11)
    normal.font.name = EN_BODY
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EA_BODY)
    normal.paragraph_format.line_spacing = 1.4
    # 内置 Heading 1-6 指定东亚字体为微软雅黑（标题样式自带大纲级别，导航窗格据此生成多级目录）
    for _hl in range(1, 7):
        try:
            doc.styles["Heading %d" % _hl].element.get_or_add_rPr(
                ).get_or_add_rFonts().set(qn("w:eastAsia"), EA_HEAD)
        except KeyError:
            pass

    i, n = 0, len(lines); in_code = False; code_buf = []; table_buf = []
    def close_table():
        nonlocal table_buf
        if table_buf:
            flush_table(doc, table_buf); table_buf = []
    while i < n:
        line = lines[i]
        if line.strip().startswith("```"):
            if not in_code:
                close_table(); in_code = True; code_buf = []
            else:
                p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.left_indent = Pt(6); p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.space_after=Pt(6); shade_par(p, "F5F5F5")
                for k, cl in enumerate(code_buf):
                    if k: p.add_run().add_break()
                    set_font(p.add_run(cl), en=EN_CODE, ea=EN_CODE, size=9.5,
                             color=RGBColor(0x22,0x22,0x22))
                in_code = False; code_buf = []
            i += 1; continue
        if in_code:
            code_buf.append(line); i += 1; continue
        if line.lstrip().startswith("|") and line.rstrip().endswith("|"):
            table_buf.append(line); i += 1; continue
        else:
            close_table()
        s = line.strip()
        if not s:
            i += 1; continue
        if s == "---" or set(s) <= set("-*_ "):
            i += 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            lv = len(m.group(1)); sizes = {1:20,2:16,3:13.5,4:12,5:11.5,6:11}
            p = doc.add_paragraph(style="Heading %d" % min(lv,6))  # 内置标题样式: 导航窗格识别多级层级
            p.paragraph_format.space_before = Pt(10 if lv<=2 else 6)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, m.group(2), size=sizes.get(lv,11), base_ea=EA_HEAD,
                       color=HEAD_COLOR, bold=True)
            i += 1; continue
        mq = re.match(r"^>\s?(.*)$", s)
        if mq:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(14)
            p.paragraph_format.line_spacing = 1.3
            add_inline(p, mq.group(1), size=10.5, color=QUOTE_COLOR)
            i += 1; continue
        mu = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        mo = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if mu or mo:
            indent = len((mu or mo).group(1).replace("\t","  "))
            lvl = indent//2
            prefix = "• " if mu else (mo.group(2)+". ")
            body = mu.group(2) if mu else mo.group(3)
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(14+14*lvl)
            p.paragraph_format.line_spacing = 1.35; p.paragraph_format.space_after = Pt(2)
            set_font(p.add_run(prefix), size=11)
            add_inline(p, body, size=11)
            i += 1; continue
        p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.45
        add_inline(p, s, size=11)
        i += 1
    close_table()
    out = os.path.splitext(md_path)[0] + ".docx"
    doc.save(out)
    return out, os.path.getsize(out)

if __name__ == "__main__":
    for md in sys.argv[1:]:
        try:
            o, sz = convert(md); print("OK  %s  (%.1f KB)" % (o, sz/1024))
        except Exception as e:
            import traceback; traceback.print_exc(); print("ERR", md, "->", repr(e))